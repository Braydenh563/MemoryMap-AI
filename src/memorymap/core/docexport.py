"""Turning a document's markdown into something to hand somebody else.

DOCUMENTS_PLAN Phase 5 item 1 and Phase 7. Two jobs live here:

* **Comments as footnotes.** A document's remarks are `==words== %%about
  them%%` in its own text (the model, and the reasons for that syntax, are in
  `frontend/documents.js` between `DOC-COMMENT-BEGIN` and `DOC-COMMENT-END`).
  Read view hides them; an export turns each one into a footnote, so a file
  handed to somebody carries the remarks rather than dropping them.
* **The other formats** a document can leave as: self-contained HTML, a .docx,
  and a markdown bundle with its images beside it (Phase 7).

**The comment conversion exists twice**, here and in documents.js, because the
browser renders Read view and the server writes the downloads and neither can
call the other. `tests/test_doc_comments.py` runs both over one fixture and
fails if their output differs by a byte, which is the only thing that keeps two
implementations of one rule honest.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

#: Where a `%%…%%` is not a comment: inside a fence or inline code it is an
#: example of the syntax, and in frontmatter it is a property's value. Same
#: three as `DOC_COMMENT_SKIP` in documents.js.
_SKIP = [
    re.compile(r"(^|\n)[ \t]*(```|~~~)[^\n]*\n[\s\S]*?(\n[ \t]*\2[^\n]*|$)"),
    re.compile(r"`[^`\n]+`"),
    re.compile(r"^---\n[\s\S]*?\n---"),
]

_COMMENT = re.compile(r"%%([^\n]*?)%%")
_TARGET = re.compile(r"==([^=\n]{1,400})==$")


@dataclass(frozen=True)
class Comment:
    """One remark, with the spans the editor's own model hands back."""

    start: int
    end: int
    body: str
    target: str
    target_start: int


def _skip_mask(text: str) -> bytearray:
    mask = bytearray(len(text))
    for pattern in _SKIP:
        for match in pattern.finditer(text):
            if match.start() == match.end():
                continue
            mask[match.start() : match.end()] = b"\x01" * (match.end() - match.start())
    return mask


def parse_comments(text: str) -> list[Comment]:
    """Every remark in the document, in document order."""
    body = text or ""
    if not body:
        return []
    mask = _skip_mask(body)
    out: list[Comment] = []
    for match in _COMMENT.finditer(body):
        start, end = match.start(), match.end()
        if mask[start]:
            continue
        note = match.group(1).strip()
        #: `%%%%` is a typo, not an empty remark.
        if not note:
            continue
        head = re.sub(r"[ \t]$", "", body[:start])
        gap = start - len(head)
        target = _TARGET.search(head)
        #: At most one space between the words and the remark about them, the
        #: same rule the editor's model states.
        anchored = target is not None and gap <= 1
        out.append(
            Comment(
                start=start,
                end=end,
                body=note,
                target=target.group(1) if anchored else "",
                target_start=(start - len(target.group(0)) - gap) if anchored else start,
            )
        )
    return out


def strip_comments(text: str) -> str:
    """The document without its remarks. The highlights stay: a highlight is
    something an author did to their own text and renders as one."""
    body = text or ""
    if not body:
        return ""
    out = body
    for comment in reversed(parse_comments(body)):
        start, end = comment.start, comment.end
        #: The space the remark was written after belongs to the remark; left
        #: behind it reads as a typo in the sentence it was about.
        if start and out[start - 1] == " ":
            start -= 1
        elif end < len(out) and out[end] == " ":
            end += 1
        line_start = out.rfind("\n", 0, max(0, start)) + 1
        line_break = out.find("\n", end)
        line_end = len(out) if line_break == -1 else line_break
        if not out[line_start:start].strip() and not out[end:line_end].strip():
            #: A line that was nothing but a comment goes whole, newline and
            #: all, or resolving the last remark leaves a paragraph break
            #: nobody wrote.
            out = out[:line_start] + out[min(line_end + 1, len(out)) :]
            continue
        out = out[:start] + out[end:]
    return out


def comments_to_footnotes(text: str, prefix: str = "c") -> str:
    """The document with its remarks as footnotes: what an export carries."""
    body = text or ""
    if not body:
        return ""
    comments = parse_comments(body)
    if not comments:
        return body
    out = body
    notes: list[str] = []
    for index in range(len(comments) - 1, -1, -1):
        comment = comments[index]
        label = f"{prefix}{index + 1}"
        notes.insert(0, f"[^{label}]: {comment.body}")
        #: The reference replaces the remark and nothing else, so the marker
        #: sits against the words it is about.
        out = f"{out[: comment.start]}[^{label}]{out[comment.end :]}"
    tail = "" if out.endswith("\n") else "\n"
    return f"{out}{tail}\n" + "\n".join(notes) + "\n"


#: Where an image goes inside the bundle. A folder rather than the root so the
#: markdown file is the first thing a person sees when they open the zip, and
#: a name that says what is in it without being clever about it.
ASSET_DIR = "assets"

_IMAGE_LINK = re.compile(r"(!?\[[^\]]*\]\()(/media/([A-Za-z0-9._-]+))(\))")


def rewrite_media_links(text: str, names: set[str]) -> str:
    """`/media/x.png` becomes `assets/x.png`, for the files that travel.

    Only for the names actually put in the bundle: a link to a file that is no
    longer on disk keeps its original address rather than pointing at a folder
    entry that does not exist, which is the difference between a bundle with a
    missing picture and one with a broken relative link in it.
    """
    if not text:
        return ""

    def swap(match: re.Match[str]) -> str:
        name = match.group(3)
        if name not in names:
            return match.group(0)
        return f"{match.group(1)}{ASSET_DIR}/{name}{match.group(4)}"

    return _IMAGE_LINK.sub(swap, text)


def bundle(title: str, text: str, media_dir: Path, names: set[str]) -> tuple[bytes, list[str]]:
    """The document and its pictures as one zip, and what went into it.

    Phase 7's "markdown with assets". The markdown a person gets out of this is
    the same markdown `export.md` gives them (title as an H1, comments as
    footnotes), with one difference: every image link that resolves to a file
    in this notebook now points at `assets/` beside it, so the bundle opens in
    any markdown reader with its pictures showing and nothing else installed.

    A name that is not on disk is skipped and its link left alone rather than
    rewritten to nothing: a bundle that lies about what is in it is worse than
    one that is honestly short of a picture. The caller decides which names are
    even candidates (the document's own `/media/` references, through
    `media_gc.referenced_names`, which is the same parse the collector uses to
    decide what is not an orphan).

    Returns the zip's bytes and the asset names it carries, in the order they
    were written, so a route or a test can say what travelled.
    """
    root = Path(media_dir)
    carried: list[str] = []
    files: list[tuple[str, bytes]] = []
    for name in sorted(names):
        #: `resolve()` and the containment check together: the names come out
        #: of a document's own text, which is user input, and `..` in one of
        #: them must not reach a file outside the media folder. The name regex
        #: already refuses a slash; this is the belt to that braces.
        path = (root / name).resolve()
        try:
            inside = path.is_relative_to(root.resolve())
        except (OSError, ValueError):
            continue
        if not inside or not path.is_file():
            continue
        try:
            files.append((name, path.read_bytes()))
        except OSError:
            continue
        carried.append(name)
    body = f"# {title}\n\n{comments_to_footnotes(text or '')}"
    body = rewrite_media_links(body, set(carried))
    buffer = io.BytesIO()
    #: Deflate, and one archive written in memory: a document with its images
    #: is a handful of megabytes at the outside, and streaming a zip to disk
    #: first would need a temporary file whose cleanup is one more thing to get
    #: wrong on a crash.
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{_stem(title)}.md", body)
        for name, data in files:
            archive.writestr(f"{ASSET_DIR}/{name}", data)
    return buffer.getvalue(), carried


def _stem(title: str) -> str:
    """The markdown file's own name inside the zip. Same shape as the route's
    `_safe_filename`, and here rather than imported from it because a core
    module must not depend on the API layer."""
    cleaned = re.sub(r"[^\w\s-]", "", title or "").strip() or "document"
    return re.sub(r"[\s_]+", "-", cleaned)[:60]


#: python-docx, if this install has it. An optional extra by decision
#: (DOCUMENTS_PLAN Phase 7): a .docx writer is a dependency most people who
#: use this app will never need, and the suite must run without it.
def docx_available() -> bool:
    try:
        import docx  # noqa: F401
    except Exception:
        return False
    return True


#: The markdown this converter understands, and it is written down so its
#: limits are too: headings, bullet and numbered lists (nested by indent, and
#: task boxes), quotes, pipe tables, fenced code, and inline bold, italic,
#: strike, code and links. Suggestion mode's marks (`{++…++}`, `{--…--}`,
#: documents.js `DOC-SUGGEST`) become Word's own tracked changes, so a draft
#: under review opens in Word with its revisions to accept or reject there.
#: Everything else (callouts, embeds, maths) stays the paragraph it was, as
#: text: a converter that half-renders them would be worse than one whose
#: limits are said.
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_TASK = re.compile(r"^\[([ xX])\]\s+(.*)$")
_QUOTE = re.compile(r"^>\s?(.*)$")
_FENCE = re.compile(r"^\s*(```|~~~)")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_INLINE = re.compile(
    r"(\{\+\+.+?\+\+\}|\{--.+?--\}|\[[^\]\n]+\]\([^)\s]+\)|\*\*\*[^*\n]+\*\*\*|\*\*[^*\n]+\*\*"
    r"|\*[^*\n]+\*|~~[^~\n]+~~|`[^`\n]+`)"
)
_LINK = re.compile(r"^\[([^\]\n]+)\]\(([^)\s]+)\)$")
#: The only addresses a link in a handed-over file may point at. A relative
#: path means something inside this notebook and nothing in Word; anything
#: else (`javascript:`, `file:`) is not something to hand to another program.
_SAFE_LINK = re.compile(r"^(https?://|mailto:)", re.IGNORECASE)
#: Who a tracked change is by, as Word shows it in the margin.
REVISION_AUTHOR = "MemoryMap"


def to_docx(title: str, text: str) -> bytes:
    """The document as a .docx. Raises RuntimeError when the extra is absent.

    The caller checks `docx_available()` first and answers 501 rather than
    500: "this install does not have the Word exporter" is a state of the
    install, not a failure of the request.
    """
    try:
        import docx
    except Exception as error:  # pragma: no cover - the guard above is the path
        raise RuntimeError("python-docx is not installed") from error

    document = docx.Document()
    document.add_heading(title or "Document", level=0)
    state = {"revision": 0}
    lines = (comments_to_footnotes(text or "")).split("\n")
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        index += 1
        if _FENCE.match(line):
            #: A fence is code, a line at a time, in a monospace face and
            #: never read as markdown: `**` in a sample is two asterisks.
            while index < len(lines) and not _FENCE.match(lines[index]):
                run = document.add_paragraph(style="No Spacing").add_run(lines[index].rstrip())
                run.font.name = "Consolas"
                index += 1
            index += 1
            continue
        if not line.strip():
            continue
        if "|" in line and index < len(lines) and _TABLE_RULE.match(lines[index]):
            rows = [_cells(line)]
            index += 1
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_cells(lines[index]))
                index += 1
            _table(document, rows, state)
            continue
        heading = _HEADING.match(line)
        if heading:
            #: Word's own levels stop at 9 and markdown's at 6, so no clamp is
            #: needed beyond the pattern itself.
            paragraph = document.add_heading("", level=len(heading.group(1)))
            _inline(paragraph, heading.group(2), state)
            continue
        listed = _BULLET.match(line) or _NUMBERED.match(line)
        if listed:
            depth = min(2, len(listed.group(1).replace("\t", "    ")) // 2)
            kind = "List Bullet" if _BULLET.match(line) else "List Number"
            body = listed.group(2)
            task = _TASK.match(body)
            if task:
                body = ("☑ " if task.group(1) in "xX" else "☐ ") + task.group(2)
            paragraph = document.add_paragraph(style=kind if depth == 0 else f"{kind} {depth + 1}")
            _inline(paragraph, body, state)
            continue
        quoted = _QUOTE.match(line)
        if quoted:
            _inline(document.add_paragraph(style="Intense Quote"), quoted.group(1), state)
            continue
        _inline(document.add_paragraph(), line, state)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _cells(line: str) -> list[str]:
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    return [cell.strip() for cell in body.split("|")]


def _table(document, rows: list[list[str]], state: dict) -> None:
    """A pipe table as a Word table, its first row the header, in bold."""
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            paragraph = table.cell(r, c).paragraphs[0]
            _inline(paragraph, row[c] if c < len(row) else "", state, bold=r == 0)


def _plain(text: str) -> str:
    """Markdown's emphasis markers off a run of text."""
    return re.sub(r"\*\*|\*|`", "", text)


def _inline(paragraph, text: str, state: dict, bold: bool = False, holder=None, deleted: bool = False) -> None:
    """One line's inline markdown as Word runs, appended to `paragraph`, or to
    `holder` (a tracked change or a link) when given."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    def add(piece: str, **marks):
        run = paragraph.add_run(piece)
        run.bold = marks.get("bold") or bold or None
        run.italic = marks.get("italic") or None
        if marks.get("strike"):
            run.font.strike = True
        if marks.get("code"):
            run.font.name = "Consolas"
        if marks.get("link"):
            run.font.underline = True
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        if deleted:
            #: Deleted text is `w:delText` in a revision, not `w:t`: Word
            #: shows it struck in the margin rather than as live text.
            for node in run._r.findall(qn("w:t")):
                node.tag = qn("w:delText")
        if holder is not None:
            holder.append(run._r)
        return run

    for piece in _INLINE.split(text or ""):
        if not piece:
            continue
        if piece.startswith(("{++", "{--")) and piece.endswith(("++}", "--}")) and len(piece) > 6:
            state["revision"] += 1
            revision = OxmlElement("w:ins" if piece.startswith("{++") else "w:del")
            revision.set(qn("w:id"), str(state["revision"]))
            revision.set(qn("w:author"), REVISION_AUTHOR)
            (holder if holder is not None else paragraph._p).append(revision)
            _inline(paragraph, piece[3:-3], state, bold=bold, holder=revision, deleted=piece.startswith("{--"))
            continue
        link = _LINK.match(piece)
        if link and _SAFE_LINK.match(link.group(2)) and holder is None:
            from docx.opc.constants import RELATIONSHIP_TYPE

            element = OxmlElement("w:hyperlink")
            element.set(
                qn("r:id"),
                paragraph.part.relate_to(link.group(2), RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
            )
            (holder if holder is not None else paragraph._p).append(element)
            run = paragraph.add_run(_plain(link.group(1)))
            run.bold = bold or None
            run.font.underline = True
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            element.append(run._r)
            continue
        if link:
            add(_plain(link.group(1)))
        elif piece.startswith("***") and piece.endswith("***") and len(piece) > 6:
            add(piece[3:-3], bold=True, italic=True)
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add(piece[2:-2], bold=True)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            add(piece[1:-1], italic=True)
        elif piece.startswith("~~") and piece.endswith("~~") and len(piece) > 4:
            add(piece[2:-2], strike=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            add(piece[1:-1], code=True)
        else:
            add(piece)
